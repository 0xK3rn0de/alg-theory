from docx import Document

def save_report_to_docx(data, filename="report.docx"):
    """
    Сохраняет отчёт в формате .docx.
    data: список словарей с ключами 'name', 'size', 'fabric', 'cost'
    """
    doc = Document()
    doc.add_heading('Отчёт по расчёту одежды', 0)

    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Изделие'
    hdr_cells[1].text = 'Размер'
    hdr_cells[2].text = 'Расход ткани (м)'
    hdr_cells[3].text = 'Стоимость (руб.)'

    for item in data:
        row_cells = table.add_row().cells
        row_cells[0].text = item['name']
        row_cells[1].text = str(item['size'])
        row_cells[2].text = f"{item['fabric']:.2f}"
        row_cells[3].text = f"{item['cost']:.2f}"

    doc.save(filename)
    print(f"Отчёт сохранён в файл {filename}")